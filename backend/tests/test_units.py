import io
from itertools import pairwise

import docx
import pytest

from app.documents import (
    ExtractionError,
    _clean,
    chunk_pages,
    extract_pages,
    split_text,
)
from app.rag import build_contents, make_title, remove_invalid_citations


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal valid PDF with one line of text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_split_text_respects_size_and_covers_everything():
    text = " ".join(f"Sentence number {i} is here." for i in range(400))
    pieces = split_text(text, size=300, overlap=60)
    assert len(pieces) > 1
    assert all(len(piece) <= 300 for _, piece in pieces)
    # Every word of the input appears in some chunk.
    joined = " ".join(piece for _, piece in pieces)
    assert all(f"number {i} " in joined for i in range(400))
    # Pieces overlap: each starts before the previous one ended.
    for (start_a, piece_a), (start_b, _) in pairwise(pieces):
        assert start_b < start_a + len(piece_a)


def test_split_text_handles_text_without_spaces():
    pieces = split_text("x" * 1000, size=300, overlap=50)
    assert "".join(piece for _, piece in pieces).count("x") >= 1000
    assert all(len(piece) <= 300 for _, piece in pieces)


def test_short_text_is_one_chunk():
    assert split_text("Hello world.", size=300, overlap=50) == [(0, "Hello world.")]
    assert split_text("", size=300, overlap=50) == []


def test_chunk_pages_tracks_pages():
    pages = [(1, "alpha " * 100), (2, ""), (3, "omega " * 100)]
    chunks = chunk_pages(pages, size=200, overlap=20)
    assert chunks[0][0] == 1
    assert chunks[-1][0] == 3
    assert {page for page, _ in chunks} == {1, 3}
    assert all("omega" not in text for page, text in chunks if page == 1)
    assert all("alpha" not in text for page, text in chunks if page == 3)


def test_clean_normalises_whitespace():
    assert _clean("a \t b\r\n\n\n\nc\x00 ") == "a b\n\nc"


def test_extract_text_and_markdown():
    assert extract_pages("notes.md", "# Title\nCafé".encode()) == [(None, "# Title\nCafé")]
    assert extract_pages("legacy.txt", "Caf\xe9".encode("cp1252")) == [(None, "Café")]


def test_extract_csv_rows():
    [(page, text)] = extract_pages("data.csv", b"name,dept\nAda,Research\nBob,\n")
    assert page is None
    assert text == "name: Ada; dept: Research\nname: Bob"


def test_extract_docx():
    document = docx.Document()
    document.add_paragraph("Hello from Word.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"
    buffer = io.BytesIO()
    document.save(buffer)
    [(_, text)] = extract_pages("file.docx", buffer.getvalue())
    assert "Hello from Word." in text and "Cell A | Cell B" in text


def test_extract_pdf_pages():
    pages = extract_pages("file.pdf", make_pdf(["First page text", "Second page text"]))
    assert [page for page, _ in pages] == [1, 2]
    assert "First page text" in pages[0][1]
    assert "Second page text" in pages[1][1]


def test_broken_files_raise_readable_errors():
    with pytest.raises(ExtractionError, match="PDF"):
        extract_pages("broken.pdf", b"not a pdf")
    with pytest.raises(ExtractionError, match="Word"):
        extract_pages("broken.docx", b"not a docx")


def test_remove_invalid_citations():
    text = "A [1], B [2][3], C [4], D [0], E [12]."
    assert remove_invalid_citations(text, 3) == "A [1], B [2][3], C, D, E."


def test_build_contents_alternates_and_starts_with_user():
    history = [("assistant", "old answer"), ("user", "q1"), ("user", "q2"), ("assistant", "")]
    contents = build_contents(history, "prompt")
    assert [turn["role"] for turn in contents] == ["user"]
    assert contents[0]["parts"][0]["text"] == "q1\n\nq2\n\nprompt"

    contents = build_contents([("user", "q1"), ("assistant", "a1")], "prompt")
    assert [turn["role"] for turn in contents] == ["user", "model", "user"]


def test_make_title():
    assert make_title("Short  question\n here?") == "Short question here?"
    long = "How much parental leave do primary caregivers get, and when are reports due?"
    assert make_title(long) == "How much parental leave do primary caregivers get, and when…"
    assert len(make_title("x" * 100)) == 61
