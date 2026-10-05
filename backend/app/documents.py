"""Turning uploaded files into embedded chunks."""

import asyncio
import csv
import io
import logging
import re
from pathlib import Path

import docx
import numpy as np
from pypdf import PdfReader

from app import gemini
from app.config import settings
from app.db import connect

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}

# A list of (page number or None, text) pairs.
Pages = list[tuple[int | None, str]]


class ExtractionError(Exception):
    """A file couldn't be read. The message is safe to show to the user."""


# -- extraction ---------------------------------------------------------------


def extract_pages(filename: str, data: bytes) -> Pages:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return _pdf_pages(data)
    if extension == ".docx":
        return [(None, _docx_text(data))]
    if extension == ".csv":
        return [(None, _csv_text(_decode(data)))]
    return [(None, _decode(data))]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf_pages(data: bytes) -> Pages:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("This PDF is password-protected. Remove the password first.")
        return [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, 1)]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("This PDF couldn't be read. It may be damaged.") from exc


def _docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This Word file couldn't be read. It may be damaged.") from exc
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _csv_text(text: str) -> str:
    """Render each row as "column: value; ..." so rows read well on their own."""
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error:
        return text
    if len(rows) < 2:
        return text
    header, *body = rows
    lines = []
    for row in body:
        cells = [
            f"{name.strip()}: {value.strip()}" for name, value in zip(header, row, strict=False)
        ]
        lines.append("; ".join(cell for cell in cells if not cell.endswith(": ")))
    return "\n".join(line for line in lines if line)


# -- chunking -----------------------------------------------------------------


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_text(text: str, size: int, overlap: int) -> list[tuple[int, str]]:
    """Split text into overlapping pieces of at most `size` characters.

    Each piece ends at a paragraph, line, sentence or word break where possible.
    Returns (start offset, piece) pairs.
    """
    overlap = max(0, min(overlap, size // 4))
    pieces: list[tuple[int, str]] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            for separator in ("\n\n", "\n", ". ", " "):
                cut = text.rfind(separator, start + size // 2, end)
                if cut != -1:
                    end = cut + len(separator)
                    break
        piece = text[start:end].strip()
        if piece:
            pieces.append((start, piece))
        if end >= len(text):
            break
        # Step back by the overlap, then forward to the next word start.
        next_start = max(end - overlap, start + 1)
        space = text.find(" ", next_start, end)
        start = space + 1 if space != -1 else next_start
    return pieces


def chunk_pages(pages: Pages, size: int, overlap: int) -> Pages:
    """Chunk each page separately, so every chunk has an exact page number."""
    return [
        (page, piece)
        for page, text in pages
        for _, piece in split_text(_clean(text), size, overlap)
    ]


# -- ingestion ----------------------------------------------------------------


def to_blob(vector: list[float]) -> bytes:
    array = np.asarray(vector, dtype=np.float32)
    norm = float(np.linalg.norm(array))
    return (array / norm if norm else array).tobytes()


async def ingest(document_id: str, filename: str, data: bytes) -> None:
    """Extract, chunk, embed and store a document. Runs in the background."""
    try:
        pages = await asyncio.to_thread(extract_pages, filename, data)
        chunks = chunk_pages(pages, settings.CHUNK_SIZE, settings.CHUNK_OVERLAP)
        if not chunks:
            raise ExtractionError(
                "No readable text was found in this file. Scanned PDFs need OCR first."
            )
        vectors = await gemini.embed([text for _, text in chunks], "RETRIEVAL_DOCUMENT")
        with connect() as conn:
            if conn.execute("SELECT 1 FROM documents WHERE id = ?", (document_id,)).fetchone():
                conn.executemany(
                    "INSERT INTO chunks (document_id, position, page, text, embedding)"
                    " VALUES (?, ?, ?, ?, ?)",
                    [
                        (document_id, position, page, text, to_blob(vector))
                        for position, ((page, text), vector) in enumerate(
                            zip(chunks, vectors, strict=True)
                        )
                    ],
                )
                conn.execute(
                    "UPDATE documents SET status = 'ready', chunk_count = ? WHERE id = ?",
                    (len(chunks), document_id),
                )
        logger.info("Indexed %s into %d chunks", filename, len(chunks))
    except Exception as exc:
        if isinstance(exc, (ExtractionError, gemini.GeminiError)):
            error = str(exc)
        else:
            logger.exception("Failed to process %s", filename)
            error = "Something went wrong while processing this file."
        with connect() as conn:
            conn.execute(
                "UPDATE documents SET status = 'failed', error = ? WHERE id = ?",
                (error, document_id),
            )
